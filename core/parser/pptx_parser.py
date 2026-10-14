"""PPTX-Parser mit python-pptx.

Extrahiert Folieninhalt und Notizen. Jede Folie wird ein eigener Abschnitt.
"""

from pathlib import Path

from pptx import Presentation

from core.logging_config import get_logger
from core.parser.base import DocumentParser, ParsedDocument, Section

logger = get_logger("parser.pptx")


class PptxParser(DocumentParser):
    """Parser für PowerPoint-Dateien (.pptx)."""

    def supported_extensions(self) -> set[str]:
        return {".pptx"}

    def parse(self, file_path: Path) -> ParsedDocument:
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Datei nicht gefunden: {file_path}")

        try:
            prs = Presentation(str(file_path))
        except Exception as e:
            raise ValueError(f"PPTX kann nicht geöffnet werden: {file_path} ({e})") from e

        sections: list[Section] = []
        all_text_parts: list[str] = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_title = self._get_slide_title(slide)
            slide_text = self._extract_slide_text(slide)
            notes_text = self._extract_notes(slide)

            content_parts: list[str] = []
            if slide_text:
                content_parts.append(slide_text)
            if notes_text:
                content_parts.append(f"[Notizen] {notes_text}")

            content = "\n".join(content_parts)
            if content.strip():
                title = slide_title or f"Folie {slide_num}"
                sections.append(
                    Section(
                        title=title,
                        content=content,
                        level=1,
                        page_number=slide_num,
                    )
                )
                all_text_parts.append(f"--- Folie {slide_num}: {title} ---\n{content}")

        properties = self._extract_properties(prs)
        full_text = "\n\n".join(all_text_parts)

        logger.info(
            "PPTX geparst: %s (%d Folien, %d Zeichen)",
            file_path.name,
            len(sections),
            len(full_text),
        )

        return ParsedDocument(
            file_path=str(file_path),
            file_name=file_path.name,
            text=full_text,
            sections=sections,
            properties=properties,
        )

    def _get_slide_title(self, slide) -> str:
        """Extrahiert den Folientitel (falls vorhanden)."""
        if slide.shapes.title and slide.shapes.title.text:
            return slide.shapes.title.text.strip()
        return ""

    def _extract_slide_text(self, slide) -> str:
        """Extrahiert den gesamten Text einer Folie."""
        text_parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        text_parts.append(text)
            if shape.has_table:
                table_text = self._extract_table(shape.table)
                if table_text:
                    text_parts.append(table_text)
        return "\n".join(text_parts)

    def _extract_table(self, table) -> str:
        """Extrahiert eine Tabelle als Text."""
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        return "\n".join(rows)

    def _extract_notes(self, slide) -> str:
        """Extrahiert die Notizen einer Folie."""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            text = slide.notes_slide.notes_text_frame.text.strip()
            return text
        return ""

    def _extract_properties(self, prs: Presentation) -> dict[str, str]:
        """Liest Dokument-Properties aus der PPTX-Datei."""
        properties: dict[str, str] = {}
        props = prs.core_properties
        if props.title:
            properties["title"] = props.title
        if props.author:
            properties["author"] = props.author
        if props.subject:
            properties["subject"] = props.subject
        if props.created:
            properties["creation_date"] = props.created.isoformat()
        if props.modified:
            properties["modification_date"] = props.modified.isoformat()
        if props.last_modified_by:
            properties["last_modified_by"] = props.last_modified_by
        return properties
