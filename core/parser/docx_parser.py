"""DOCX-Parser mit python-docx.

Extrahiert Text mit Strukturinformationen (Überschriften-Styles)
und Dokument-Properties.
"""

from pathlib import Path

import docx

from core.logging_config import get_logger
from core.parser.base import DocumentParser, ParsedDocument, Section

logger = get_logger("parser.docx")

# Mapping von Word-Heading-Styles auf Level
_HEADING_STYLES = {
    "Heading 1": 1, "Heading 2": 2, "Heading 3": 3,
    "Heading 4": 4, "Heading 5": 5, "Heading 6": 6,
    "Überschrift 1": 1, "Überschrift 2": 2, "Überschrift 3": 3,
    "Überschrift 4": 4, "Überschrift 5": 5, "Überschrift 6": 6,
}


class DocxParser(DocumentParser):
    """Parser für Word-Dokumente (.docx)."""

    def supported_extensions(self) -> set[str]:
        return {".docx"}

    def parse(self, file_path: Path) -> ParsedDocument:
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Datei nicht gefunden: {file_path}")

        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise ValueError(f"DOCX kann nicht geöffnet werden: {file_path} ({e})") from e

        sections = self._extract_sections(doc)
        properties = self._extract_properties(doc)
        full_text = "\n\n".join(
            f"{s.title}\n{s.content}" if s.level > 0 else s.content
            for s in sections
        )

        logger.info(
            "DOCX geparst: %s (%d Abschnitte, %d Zeichen)",
            file_path.name,
            len(sections),
            len(full_text),
        )

        return ParsedDocument(
            file_path=str(file_path),
            file_name=file_path.name,
            text=full_text,
            sections=sections,
            properties=properties,
        )

    def _extract_sections(self, doc: docx.Document) -> list[Section]:
        """Extrahiert Abschnitte basierend auf Heading-Styles."""
        sections: list[Section] = []
        current_title = ""
        current_level = 0
        current_paragraphs: list[str] = []

        for para in doc.paragraphs:
            style_name = para.style.name if para.style else ""
            heading_level = self._get_heading_level(style_name)

            if heading_level > 0:
                # Vorherigen Abschnitt speichern
                if current_paragraphs:
                    content = "\n".join(current_paragraphs).strip()
                    if content:
                        sections.append(
                            Section(
                                title=current_title or "Einleitung",
                                content=content,
                                level=current_level,
                            )
                        )
                current_title = para.text.strip()
                current_level = heading_level
                current_paragraphs = []
            else:
                text = para.text.strip()
                if text:
                    current_paragraphs.append(text)

        # Letzten Abschnitt speichern
        if current_paragraphs:
            content = "\n".join(current_paragraphs).strip()
            if content:
                sections.append(
                    Section(
                        title=current_title or "Inhalt",
                        content=content,
                        level=current_level,
                    )
                )

        # Tabellen extrahieren und als eigene Abschnitte hinzufügen
        table_sections = self._extract_tables(doc)
        if table_sections:
            sections.extend(table_sections)

        return sections if sections else [
            Section(title="Inhalt", content=self._fallback_text(doc), level=0)
        ]

    def _get_heading_level(self, style_name: str) -> int:
        """Bestimmt das Heading-Level aus dem Style-Namen."""
        if style_name in _HEADING_STYLES:
            return _HEADING_STYLES[style_name]
        # Auch "Heading1" ohne Leerzeichen prüfen
        for key, level in _HEADING_STYLES.items():
            if style_name.lower().replace(" ", "") == key.lower().replace(" ", ""):
                return level
        return 0

    def _extract_tables(self, doc: docx.Document) -> list[Section]:
        """Extrahiert Tabellen als strukturierten Text."""
        sections: list[Section] = []
        for i, table in enumerate(doc.tables):
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                content = "\n".join(rows)
                sections.append(
                    Section(
                        title=f"Tabelle {i + 1}",
                        content=content,
                        level=0,
                    )
                )
        return sections

    def _extract_properties(self, doc: docx.Document) -> dict[str, str]:
        """Liest Dokument-Properties aus der DOCX-Datei."""
        properties: dict[str, str] = {}
        props = doc.core_properties
        if props.title:
            properties["title"] = props.title
        if props.author:
            properties["author"] = props.author
        if props.subject:
            properties["subject"] = props.subject
        if props.created:
            properties["creation_date"] = props.created.isoformat()
        if props.modified:
            properties["modification_date"] = props.modified.isoformat()
        if props.last_modified_by:
            properties["last_modified_by"] = props.last_modified_by
        if props.revision:
            properties["revision"] = str(props.revision)
        if props.category:
            properties["category"] = props.category
        if props.keywords:
            properties["keywords"] = props.keywords
        return properties

    def _fallback_text(self, doc: docx.Document) -> str:
        """Fallback: Alle Absätze als Fließtext extrahieren."""
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
